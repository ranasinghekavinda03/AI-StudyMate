"""Text extraction helpers for supported lecture document formats."""

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError


class DocumentExtractionError(ValueError):
    """Raised when a document cannot be parsed or contains no useful text."""


@dataclass(frozen=True)
class ExtractedPage:
    """A text unit with a real source page number when the format provides one."""

    text: str
    page_number: int | None


@dataclass(frozen=True)
class ExtractedDocument:
    pages: list[ExtractedPage]

    @property
    def text(self) -> str:
        return "\n\n".join(page.text for page in self.pages if page.text).strip()

    @property
    def page_count(self) -> int:
        numbered_pages = [page.page_number for page in self.pages if page.page_number is not None]
        return max(numbered_pages, default=0)


def extract_document(data: bytes, filename: str) -> ExtractedDocument:
    """Extract text from PDF, DOCX, or UTF-8 TXT bytes."""
    extension = Path(filename).suffix.lower()

    try:
        if extension == ".pdf":
            reader = PdfReader(BytesIO(data))
            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception as exc:
                    raise DocumentExtractionError("Password-protected PDFs are not supported.") from exc
            pages = [
                ExtractedPage(text=(page.extract_text() or "").strip(), page_number=index)
                for index, page in enumerate(reader.pages, start=1)
            ]
        elif extension == ".docx":
            document = Document(BytesIO(data))
            blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        blocks.append("\t".join(cells))
            pages = [ExtractedPage(text="\n".join(blocks), page_number=None)]
        elif extension == ".txt":
            pages = [ExtractedPage(text=data.decode("utf-8-sig").strip(), page_number=None)]
        else:
            raise DocumentExtractionError("Unsupported file type. Upload a PDF, DOCX, or TXT file.")
    except DocumentExtractionError:
        raise
    except UnicodeDecodeError as exc:
        raise DocumentExtractionError("TXT files must use UTF-8 encoding.") from exc
    except (PdfReadError, PackageNotFoundError, KeyError, ValueError) as exc:
        raise DocumentExtractionError(f"The {extension[1:].upper()} file is invalid or corrupted.") from exc
    except Exception as exc:
        raise DocumentExtractionError("The document could not be processed.") from exc

    extracted = ExtractedDocument(pages=pages)
    if not extracted.text:
        raise DocumentExtractionError("No extractable text was found in the document.")
    return extracted
