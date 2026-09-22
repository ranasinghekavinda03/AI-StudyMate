from io import BytesIO

from docx import Document

from app.core.config import settings
from app.models.chunk import DocumentChunk


def _create_module(client, headers):
    response = client.post(
        "/api/v1/modules",
        headers=headers,
        json={"title": "Document Processing", "code": "CS402"},
    )
    assert response.status_code == 201
    return response.json()["id"]


def test_upload_txt_extracts_text_and_creates_lecture(client, auth_headers, db_session, tmp_path, monkeypatch):
    headers, _ = auth_headers
    module_id = _create_module(client, headers)
    monkeypatch.setattr(settings, "UPLOAD_DIR", str(tmp_path))

    response = client.post(
        "/api/v1/lectures/upload",
        headers=headers,
        data={"module_id": module_id},
        files={"file": ("lesson.txt", b"Retrieval augmented generation", "text/plain")},
    )

    assert response.status_code == 201
    body = response.json()
    assert body["title"] == "lesson"
    assert body["file_type"] == "txt"
    assert body["page_count"] == 0
    assert body["chunks_count"] == 1
    assert body["extracted_text"] == "Retrieval augmented generation"
    assert len(list(tmp_path.glob("*.txt"))) == 1
    chunks = db_session.query(DocumentChunk).filter(DocumentChunk.lecture_id == body["id"]).all()
    assert len(chunks) == 1
    assert chunks[0].page_number is None
    assert chunks[0].chunk_index == 0


def test_upload_docx_extracts_paragraphs_and_tables(client, auth_headers, db_session, tmp_path, monkeypatch):
    headers, _ = auth_headers
    module_id = _create_module(client, headers)
    monkeypatch.setattr(settings, "UPLOAD_DIR", str(tmp_path))
    document = Document()
    document.add_paragraph("Neural networks learn representations.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Term"
    table.cell(0, 1).text = "Definition"
    buffer = BytesIO()
    document.save(buffer)

    response = client.post(
        "/api/v1/lectures/upload",
        headers=headers,
        data={"module_id": module_id, "title": "Neural Networks"},
        files={
            "file": (
                "notes.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 201
    body = response.json()
    assert "Neural networks learn representations." in body["extracted_text"]
    assert "Term\tDefinition" in body["extracted_text"]
    chunks = db_session.query(DocumentChunk).filter(DocumentChunk.lecture_id == body["id"]).all()
    assert chunks
    assert all(chunk.page_number is None for chunk in chunks)


def test_upload_rejects_unsupported_and_empty_files(client, auth_headers, tmp_path, monkeypatch):
    headers, _ = auth_headers
    module_id = _create_module(client, headers)
    monkeypatch.setattr(settings, "UPLOAD_DIR", str(tmp_path))

    unsupported = client.post(
        "/api/v1/lectures/upload",
        headers=headers,
        data={"module_id": module_id},
        files={"file": ("image.png", b"png", "image/png")},
    )
    assert unsupported.status_code == 415

    empty = client.post(
        "/api/v1/lectures/upload",
        headers=headers,
        data={"module_id": module_id},
        files={"file": ("empty.txt", b"", "text/plain")},
    )
    assert empty.status_code == 422
    assert not list(tmp_path.iterdir())


def test_upload_requires_owned_module(client, auth_headers, tmp_path, monkeypatch):
    headers, _ = auth_headers
    monkeypatch.setattr(settings, "UPLOAD_DIR", str(tmp_path))
    response = client.post(
        "/api/v1/lectures/upload",
        headers=headers,
        data={"module_id": "not-owned"},
        files={"file": ("lesson.txt", b"content", "text/plain")},
    )
    assert response.status_code == 404
